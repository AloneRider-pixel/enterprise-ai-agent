"""
Document management API routes.
Handles document upload, ingestion, and lifecycle management.
"""
import hashlib
import logging
import uuid
from datetime import datetime
from io import BytesIO
from typing import Optional

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile, status
from sqlalchemy import select, func
from sqlalchemy.ext.asyncio import AsyncSession

from app.auth.dependencies import get_current_user
from app.database import get_db, vector_store
from app.models.database import Document
from app.models.schemas import DocumentListItem, DocumentStatus, DocumentUploadResponse
from app.rag.chunking import text_chunker
from app.rag.embeddings import embedding_service

logger = logging.getLogger(__name__)

router = APIRouter(prefix="/api/documents", tags=["Documents"])

ALLOWED_TYPES = {
    "application/pdf": "pdf",
    "text/plain": "txt",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    "text/markdown": "md",
}


@router.post("/upload", response_model=DocumentUploadResponse)
async def upload_document(
    file: UploadFile = File(...),
    current_user: dict = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """
    Upload a document for processing.
    Supports PDF, TXT, DOCX, and Markdown files.
    """
    # Validate file type
    content_type = file.content_type
    if content_type not in ALLOWED_TYPES:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Unsupported file type: {content_type}. Allowed: {list(ALLOWED_TYPES.keys())}",
        )
    
    # Read file content
    content = await file.read()
    file_size = len(content)
    
    if file_size > 50 * 1024 * 1024:  # 50MB limit
        raise HTTPException(
            status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
            detail="File too large. Maximum size is 50MB.",
        )
    
    # Compute content hash for deduplication
    content_hash = hashlib.sha256(content).hexdigest()
    
    # Check for duplicate
    result = await db.execute(
        select(Document).where(Document.content_hash == content_hash)
    )
    existing = result.scalar_one_or_none()
    if existing:
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail=f"Document already exists: {existing.filename} (ID: {existing.id})",
        )
    
    # Extract text content
    file_type = ALLOWED_TYPES[content_type]
    text_content = _extract_text(content, file_type)
    
    if not text_content or len(text_content.strip()) < 50:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Could not extract meaningful text from the file.",
        )
    
    # Create document record
    doc_id = uuid.uuid4()
    document = Document(
        id=doc_id,
        filename=file.filename or "unknown",
        file_type=file_type,
        file_size=file_size,
        content_hash=content_hash,
        status="uploaded",
        raw_content=text_content,
        uploaded_by=uuid.UUID(current_user["user_id"]) if len(current_user["user_id"]) == 36 else None,
    )
    
    db.add(document)
    await db.commit()
    await db.refresh(document)
    
    logger.info(f"Document uploaded: {file.filename} ({file_size} bytes)")
    
    return DocumentUploadResponse(
        document_id=str(document.id),
        filename=document.filename,
        size_bytes=file_size,
        status=document.status,
        created_at=document.created_at,
    )


@router.post("/ingest/{document_id}")
async def ingest_document(
    document_id: str,
    current_user: dict = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """
    Trigger document ingestion (chunking + embedding + vector storage).
    """
    # Get document
    result = await db.execute(
        select(Document).where(Document.id == uuid.UUID(document_id))
    )
    document = result.scalar_one_or_none()
    
    if not document:
        raise HTTPException(status_code=404, detail="Document not found")
    
    if document.status == "indexed":
        return {"status": "already_indexed", "chunks": document.chunk_count}
    
    # Update status to processing
    document.status = "processing"
    await db.commit()
    
    try:
        # Chunk the text
        chunks = text_chunker.chunk_text(
            document.raw_content,
            metadata={
                "document_id": str(document.id),
                "filename": document.filename,
                "file_type": document.file_type,
            },
        )
        
        if not chunks:
            document.status = "failed"
            await db.commit()
            return {"status": "failed", "error": "No chunks generated"}
        
        # Generate embeddings in batch
        texts = [chunk["content"] for chunk in chunks]
        embeddings = await embedding_service.embed_batch(texts)
        
        # Store in pgvector
        for chunk, embedding in zip(chunks, embeddings):
            await vector_store.store_embedding(
                chunk_id=chunk["chunk_id"],
                document_id=str(document.id),
                content=chunk["content"],
                embedding=embedding,
                metadata=chunk["metadata"],
            )
        
        # Update document status
        document.status = "indexed"
        document.chunk_count = len(chunks)
        document.indexed_at = datetime.utcnow()
        await db.commit()
        
        logger.info(f"Document indexed: {document.filename} -> {len(chunks)} chunks")
        
        return {
            "status": "indexed",
            "document_id": str(document.id),
            "chunks": len(chunks),
            "total_tokens": sum(c["metadata"].get("token_count", 0) for c in chunks),
        }
    
    except Exception as e:
        logger.error(f"Ingestion failed: {e}")
        document.status = "failed"
        await db.commit()
        raise HTTPException(status_code=500, detail=f"Ingestion failed: {str(e)}")


@router.get("", response_model=DocumentListItem)
async def list_documents(
    current_user: dict = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """List all uploaded documents with their status."""
    result = await db.execute(
        select(Document).order_by(Document.created_at.desc())
    )
    documents = result.scalars().all()
    
    items = [
        DocumentStatus(
            document_id=str(doc.id),
            filename=doc.filename,
            status=doc.status,
            chunk_count=doc.chunk_count,
            created_at=doc.created_at,
            indexed_at=doc.indexed_at,
        )
        for doc in documents
    ]
    
    return DocumentListItem(documents=items, total=len(items))


@router.delete("/{document_id}")
async def delete_document(
    document_id: str,
    current_user: dict = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """Delete a document and its associated chunks."""
    # Delete vector chunks
    await vector_store.delete_document_chunks(document_id)
    
    # Delete document record
    result = await db.execute(
        select(Document).where(Document.id == uuid.UUID(document_id))
    )
    document = result.scalar_one_or_none()
    
    if not document:
        raise HTTPException(status_code=404, detail="Document not found")
    
    await db.delete(document)
    await db.commit()
    
    logger.info(f"Document deleted: {document.filename}")
    return {"status": "deleted", "document_id": document_id}


def _extract_text(content: bytes, file_type: str) -> str:
    """Extract text content from uploaded file."""
    if file_type == "txt" or file_type == "md":
        return content.decode("utf-8", errors="ignore")
    
    elif file_type == "pdf":
        try:
            from pypdf import PdfReader
            reader = PdfReader(BytesIO(content))
            text_parts = []
            for page in reader.pages:
                text = page.extract_text()
                if text:
                    text_parts.append(text)
            return "\n\n".join(text_parts)
        except Exception as e:
            logger.error(f"PDF extraction failed: {e}")
            return ""
    
    elif file_type == "docx":
        try:
            from docx import Document as DocxDocument
            doc = DocxDocument(BytesIO(content))
            text_parts = []
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    text_parts.append(paragraph.text)
            return "\n\n".join(text_parts)
        except Exception as e:
            logger.error(f"DOCX extraction failed: {e}")
            return ""
    
    return ""
